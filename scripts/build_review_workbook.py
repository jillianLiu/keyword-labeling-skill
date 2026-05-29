#!/usr/bin/env python3
"""Build an Excel review workbook with filters on every sheet."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def read_csv(path: Path) -> tuple[list[str], list[list[str]]]:
    with path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        rows = list(reader)
    if not rows:
        return [], []
    return rows[0], rows[1:]


def safe_sheet_name(name: str) -> str:
    return name[:31].replace("/", "-").replace("\\", "-").replace("*", "-").replace("?", "-")


def add_sheet(wb: Workbook, title: str, csv_path: Path, table_name: str) -> None:
    headers, rows = read_csv(csv_path)
    ws = wb.create_sheet(safe_sheet_name(title))
    if not headers:
        ws.append(["空表"])
        return

    ws.append(headers)
    for row in rows:
        ws.append(row)

    max_row = ws.max_row
    max_col = ws.max_column
    ref = f"A1:{get_column_letter(max_col)}{max_row}"
    ws.auto_filter.ref = ref
    ws.freeze_panes = "A2"

    table = Table(displayName=table_name, ref=ref)
    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    table.tableStyleInfo = style
    ws.add_table(table)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for column_cells in ws.columns:
        header = str(column_cells[0].value or "")
        max_len = len(header)
        for cell in column_cells[1:200]:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value), 60))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        width = min(max(max_len + 2, 10), 48)
        ws.column_dimensions[column_cells[0].column_letter].width = width


def main() -> int:
    parser = argparse.ArgumentParser(description="Build filtered Excel review workbook from CSV files.")
    parser.add_argument("--detail", type=Path, required=True, help="Chinese keyword detail CSV")
    parser.add_argument("--qa-summary", type=Path)
    parser.add_argument("--coverage", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    add_sheet(wb, "关键词明细表", args.detail, "KeywordDetail")
    if args.qa_summary and args.qa_summary.exists():
        add_sheet(wb, "QA汇总表", args.qa_summary, "QASummary")
    if args.coverage and args.coverage.exists():
        add_sheet(wb, "覆盖检查表", args.coverage, "CoverageReport")

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"wrote workbook to {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
