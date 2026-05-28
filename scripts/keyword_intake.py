#!/usr/bin/env python3
"""Create a keyword fact table from CSV, TSV, or XLSX keyword exports."""

from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from pathlib import Path
from typing import Iterable


FIELDNAMES = [
    "original_keyword_id",
    "keyword",
    "source_intent",
    "volume",
    "keyword_difficulty",
    "cpc",
    "near_match_type",
    "near_match_group_id",
    "near_match_group_key",
    "near_match_group_count",
    "comparison_key",
    "morphology_key",
    "action_signal",
    "object_signal",
    "format_signal",
    "modifier_signal",
    "risk_level",
    "risk_label",
    "entity_label",
    "site_fit",
    "site_mismatch_reason",
    "notes",
]

ZH_FIELDNAMES = {
    "original_keyword_id": "原始词ID",
    "keyword": "关键词",
    "source_intent": "来源意图",
    "volume": "搜索量",
    "keyword_difficulty": "关键词难度",
    "cpc": "CPC",
    "near_match_type": "近似类型",
    "near_match_group_id": "近似组ID",
    "near_match_group_key": "近似组Key",
    "near_match_group_count": "近似组数量",
    "comparison_key": "对比键",
    "morphology_key": "词形近似键",
    "action_signal": "动作信号",
    "object_signal": "对象信号",
    "format_signal": "格式信号",
    "modifier_signal": "修饰词信号",
    "risk_level": "风险等级",
    "risk_label": "风险标签",
    "entity_label": "实体标签",
    "site_fit": "站点适配度",
    "site_mismatch_reason": "不匹配原因",
    "notes": "备注",
}


ALIASES = {
    "keyword": {"keyword", "keywords", "query", "search term", "search_term"},
    "source_intent": {"intent", "source_intent", "semrush intent", "semrush_intent"},
    "volume": {"volume", "search volume", "search_volume"},
    "keyword_difficulty": {"keyword difficulty", "keyword_difficulty", "kd", "difficulty"},
    "cpc": {"cpc", "cpc (usd)", "cpc_usd"},
}


ACTION_PATTERNS = [
    ("remove-background", r"\b(remove|erase|delete)\s+(the\s+)?(background|bg)\b|\bbackground\s+remover\b"),
    ("download", r"\b(download|downloader|save)\b"),
    ("convert", r"\b(convert|converter|to mp3|to mp4|to wav)\b"),
    ("search", r"\b(search|lookup|find)\b"),
    ("login", r"\b(login|log in|sign in|signin)\b"),
    ("view", r"\b(view|viewer|watch)\b"),
    ("edit", r"\b(edit|editor|modify)\b"),
    ("generate", r"\b(generate|generator|create|make)\b"),
]

OBJECT_PATTERNS = [
    ("video", r"\b(video|videos|mp4)\b"),
    ("audio", r"\b(audio|sound|song|music|mp3|wav)\b"),
    ("image", r"\b(image|photo|picture|thumbnail|avatar|logo|icon)\b"),
    ("account", r"\b(account|profile|username|user)\b"),
    ("hashtag", r"\b(hashtag|hashtags|tag|tags)\b"),
    ("caption", r"\b(caption|captions|subtitle|subtitles)\b"),
]

FORMAT_PATTERNS = [
    ("mp3", r"\bmp3\b"),
    ("mp4", r"\bmp4\b"),
    ("wav", r"\bwav\b"),
    ("jpg/jpeg", r"\b(jpg|jpeg)\b"),
    ("png", r"\bpng\b"),
]

MODIFIER_PATTERNS = [
    ("free", r"\bfree\b"),
    ("online", r"\bonline\b"),
    ("no signup", r"\b(no signup|without signup|no sign up)\b"),
    ("no watermark", r"\b(no watermark|without watermark)\b"),
    ("best/top", r"\b(best|top)\b"),
    ("app/mobile", r"\b(app|iphone|android|ios|mobile)\b"),
    ("download", r"\b(download|downloader)\b"),
]

RISK_PATTERNS = [
    ("adult/nsfw", "high", r"\b(nsfw|porn|nude|naked|sex|xxx|onlyfans)\b"),
    ("piracy", "high", r"\b(crack|mod apk|apk mod|pirated|leak|leaked)\b"),
    ("sensitive", "medium", r"\b(ban|banned|lawsuit|court|death|dead|shooting|war)\b"),
    ("policy", "medium", r"\b(bypass|unblock|vpn|shadowban|ban appeal)\b"),
]

PLATFORM_TERMS = [
    "tiktok",
    "douyin",
    "instagram",
    "youtube",
    "facebook",
    "snapchat",
    "twitter",
    "x",
]


def norm_header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def normalize_keyword(value: str) -> str:
    value = value.strip().lower()
    value = re.sub(r"[’']", "", value)
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def morphology_key(value: str) -> str:
    words = normalize_keyword(value).split()
    converted = []
    for word in words:
        if len(word) > 4 and word.endswith("ies"):
            converted.append(word[:-3] + "y")
        elif len(word) > 4 and word.endswith("es") and not word.endswith(("ses", "xes")):
            converted.append(word[:-2])
        elif len(word) > 4 and word.endswith("s") and not word.endswith(("ss", "us")):
            converted.append(word[:-1])
        else:
            converted.append(word)
    return " ".join(converted)


def joined_matches(patterns: Iterable[tuple[str, str]], keyword: str) -> str:
    hits = []
    for label, pattern in patterns:
        if re.search(pattern, keyword, flags=re.I):
            hits.append(label)
    return ";".join(dict.fromkeys(hits)) or "none"


def risk_tags(keyword: str) -> tuple[str, str]:
    labels = []
    level = "none"
    rank = {"none": 0, "low": 1, "medium": 2, "high": 3}
    for label, candidate_level, pattern in RISK_PATTERNS:
        if re.search(pattern, keyword, flags=re.I):
            labels.append(label)
            if rank[candidate_level] > rank[level]:
                level = candidate_level
    return level, ";".join(labels) or "none"


def entity_label(keyword: str) -> str:
    normalized = normalize_keyword(keyword)
    labels = []
    for term in PLATFORM_TERMS:
        if term in normalized.split():
            labels.append(f"platform:{term}")
    return ";".join(labels) or "none"


def read_csv_or_tsv(path: Path) -> list[dict[str, object]]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open(newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f, delimiter=delimiter))


def read_xlsx(path: Path, sheet_name: str | None) -> list[dict[str, object]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return []
    names = [str(value or "").strip() for value in header]
    output = []
    for row in rows:
        if row is None or all(value is None or str(value).strip() == "" for value in row):
            continue
        output.append({names[i]: row[i] if i < len(row) else "" for i in range(len(names))})
    return output


def map_columns(rows: list[dict[str, object]]) -> dict[str, str]:
    if not rows:
        raise ValueError("input has no rows")
    headers = list(rows[0].keys())
    normalized = {norm_header(header): header for header in headers}
    mapping = {}
    for target, aliases in ALIASES.items():
        for alias in aliases:
            if alias in normalized:
                mapping[target] = normalized[alias]
                break
    if "keyword" not in mapping:
        raise ValueError("could not find keyword column")
    return mapping


def value(row: dict[str, object], column: str | None) -> str:
    if not column:
        return ""
    raw = row.get(column, "")
    return "" if raw is None else str(raw).strip()


def build_fact_rows(rows: list[dict[str, object]], prefix: str) -> list[dict[str, str]]:
    mapping = map_columns(rows)
    fact_rows = []
    for index, row in enumerate(rows, start=1):
        keyword = value(row, mapping.get("keyword"))
        comparison = normalize_keyword(keyword)
        morph = morphology_key(keyword)
        risk_level, risk_label = risk_tags(keyword)
        fact_rows.append(
            {
                "original_keyword_id": f"{prefix}-{index:05d}",
                "keyword": keyword,
                "source_intent": value(row, mapping.get("source_intent")),
                "volume": value(row, mapping.get("volume")),
                "keyword_difficulty": value(row, mapping.get("keyword_difficulty")),
                "cpc": value(row, mapping.get("cpc")),
                "near_match_type": "unique",
                "near_match_group_id": "",
                "near_match_group_key": "",
                "near_match_group_count": "1",
                "comparison_key": comparison,
                "morphology_key": morph,
                "action_signal": joined_matches(ACTION_PATTERNS, keyword),
                "object_signal": joined_matches(OBJECT_PATTERNS, keyword),
                "format_signal": joined_matches(FORMAT_PATTERNS, keyword),
                "modifier_signal": joined_matches(MODIFIER_PATTERNS, keyword),
                "risk_level": risk_level,
                "risk_label": risk_label,
                "entity_label": entity_label(keyword),
                "site_fit": "unreviewed",
                "site_mismatch_reason": "",
                "notes": "deterministic intake only; semantic annotation is independent",
            }
        )

    groups: dict[str, list[int]] = defaultdict(list)
    morph_groups: dict[str, list[int]] = defaultdict(list)
    for index, row in enumerate(fact_rows):
        groups[row["comparison_key"]].append(index)
        morph_groups[row["morphology_key"]].append(index)

    group_no = 1
    assigned: set[int] = set()
    for key, indexes in groups.items():
        if key and len(indexes) > 1:
            gid = f"NM-{group_no:05d}"
            group_no += 1
            for idx in indexes:
                fact_rows[idx]["near_match_type"] = "format_near"
                fact_rows[idx]["near_match_group_id"] = gid
                fact_rows[idx]["near_match_group_key"] = key
                fact_rows[idx]["near_match_group_count"] = str(len(indexes))
                assigned.add(idx)

    for key, indexes in morph_groups.items():
        unassigned = [idx for idx in indexes if idx not in assigned]
        if key and len(unassigned) > 1:
            gid = f"NM-{group_no:05d}"
            group_no += 1
            for idx in unassigned:
                fact_rows[idx]["near_match_type"] = "morphology_near"
                fact_rows[idx]["near_match_group_id"] = gid
                fact_rows[idx]["near_match_group_key"] = key
                fact_rows[idx]["near_match_group_count"] = str(len(unassigned))

    return fact_rows


def write_rows(path: Path, rows: list[dict[str, str]], locale: str) -> None:
    fieldnames = FIELDNAMES
    output_rows = rows
    if locale == "zh":
        fieldnames = [ZH_FIELDNAMES[field] for field in FIELDNAMES]
        output_rows = [{ZH_FIELDNAMES[key]: value for key, value in row.items()} for row in rows]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(output_rows)


def main() -> int:
    parser = argparse.ArgumentParser(description="Build keyword fact table from a raw keyword export.")
    parser.add_argument("input_file", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--sheet")
    parser.add_argument("--id-prefix", default="KW")
    parser.add_argument("--locale", choices=["en", "zh"], default="en")
    args = parser.parse_args()

    suffix = args.input_file.suffix.lower()
    if suffix == ".xlsx":
        rows = read_xlsx(args.input_file, args.sheet)
    elif suffix in {".csv", ".tsv"}:
        rows = read_csv_or_tsv(args.input_file)
    else:
        parser.error("input_file must be .xlsx, .csv, or .tsv")

    fact_rows = build_fact_rows(rows, args.id_prefix)
    write_rows(args.out, fact_rows, args.locale)
    print(f"wrote {len(fact_rows)} fact rows to {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
